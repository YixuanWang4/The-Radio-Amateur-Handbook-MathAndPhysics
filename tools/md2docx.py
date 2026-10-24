#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert Markdown (chapter or proposal) to a styled DOCX for this book.

Usage:
    python tools/md2docx.py input.md output.docx                 # A5 (default)
    python tools/md2docx.py input.md output.docx --b5            # B5, book pages
    python tools/md2docx.py input.md output.docx --a4 --cover    # A4 with cover

Design: adjustable page size (A5 / B5 / A4), 宋体 body, 黑体 accent-colored
headings, running header, three-line table with zebra stripes, shaded formula
blocks, and callout boxes built from blockquotes. Open the result in Word and
export to PDF.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

BODY_CN = "宋体"
BODY_EN = "Times New Roman"
HEAD_CN = "黑体"
HEAD_EN = "Arial"
CODE_EN = "Consolas"

# Book identity (overridable from the command line) -------------------------
BOOK_TITLE = "业余无线电实用数学物理"
BOOK_SUBTITLE = "电台里的数学与物理"
BOOK_DESCRIPTION = "写给想弄明白天线、馈线、传播与设备原理的短波爱好者"
BOOK_EDITION = "本书编写组 · 2026 年 10 月"
HEADER_TITLE = "业余无线电实用数学物理"
HEADER_AUTHOR = "本书编写组"

ACCENT = "1F3A5F"
ACCENT_RGB = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT_TINT = "E8EEF5"
ZEBRA_FILL = "F7F9FB"
CALLOUT_FILL = "F2F5F9"
FORMULA_FILL = "F4F7FA"
ADV_ACCENT = "2E6B6B"
ADV_TINT = "F2F7F6"
ADV_RGB = RGBColor(0x2E, 0x6B, 0x6B)
RULE_GRAY = "D9D9D9"
TABLE_BORDER = ACCENT

GRAY = RGBColor(0x59, 0x59, 0x59)
MID_GRAY = RGBColor(0x7F, 0x7F, 0x7F)
BLACK = RGBColor(0x22, 0x22, 0x22)

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RE = re.compile(r"^[-*]\s+")
NUMBER_RE = re.compile(r"^\d+\.\s+")
TABLE_SEP_RE = re.compile(r"^:?-{3,}:?$")
FORMULA_RE = re.compile(r"^\$\$(.+?)\$\$$")
TITLE_RE = re.compile(r"^\*\*(.+?)\*\*$")
IMAGE_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")


# ---------------------------------------------------------------------------
# OOXML helpers
# ---------------------------------------------------------------------------

def _rfonts(rpr):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    return rf


def style_fonts(style, cn, en, size=None, bold=None, color=None):
    style.font.name = en
    rpr = style.element.get_or_add_rPr()
    rf = _rfonts(rpr)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def run_fonts(run, cn, en, size=None, bold=None, color=None):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rf = _rfonts(rpr)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_paragraph(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.insert_element_before(
        shd, "w:tabs", "w:spacing", "w:ind", "w:contextualSpacing",
        "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
        "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
        "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def paragraph_borders(paragraph, edges):
    """edges: {"top"|"left"|"bottom"|"right": (color, sz_eighths, space_pt)}"""
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        if edge not in edges:
            continue
        color, sz, space = edges[edge]
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        pbdr.append(el)
    pPr.insert_element_before(
        pbdr, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:contextualSpacing",
        "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
        "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
        "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def bottom_border(paragraph, color, sz, space=1):
    paragraph_borders(paragraph, {"bottom": (color, sz, space)})


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.insert_element_before(
        shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
        "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
        "w:cellMerge", "w:tcPrChange",
    )


def cell_bottom_border(cell, color, sz, space=0):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    tcPr.insert_element_before(
        borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
        "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers", "w:cellIns",
        "w:cellDel", "w:cellMerge", "w:tcPrChange",
    )


def table_borders(table):
    """Three-line table: accent rules top/bottom, no vertical or inner rules."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag, val, sz, color in (
        ("top", "single", 12, TABLE_BORDER),
        ("left", "nil", 0, "auto"),
        ("bottom", "single", 12, TABLE_BORDER),
        ("right", "nil", 0, "auto"),
        ("insideH", "nil", 0, "auto"),
        ("insideV", "nil", 0, "auto"),
    ):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), val)
        if val != "nil":
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), color)
        el.set(qn("w:space"), "0")
        borders.append(el)
    tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription",
    )


def mark_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    trPr.append(tblHeader)


def no_split_row(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.insert_element_before(
        cant, "w:trHeight", "w:tblHeader", "w:tblCellSpacing",
        "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange",
    )


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(begin)
    r.append(instr)
    r.append(sep)
    r.append(txt)
    r.append(end)
    return run


# ---------------------------------------------------------------------------
# document setup
# ---------------------------------------------------------------------------

def text_width_cm(doc):
    sec = doc.sections[0]
    return Emu(sec.page_width - sec.left_margin - sec.right_margin).cm


def setup_page(doc, page_size):
    section = doc.sections[0]
    if page_size == "a4":
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    elif page_size == "b5":
        section.page_width = Cm(17.6)
        section.page_height = Cm(25.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
    else:
        section.page_width = Cm(14.8)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.2)
    return section


def setup_styles(doc, page_size="a5"):
    if page_size == "a4":
        body_size, body_spacing = 10.5, 1.55
        specs = {
            "Heading 1": (17, 12, 8, 1.2),
            "Heading 2": (13, 10, 6, 1.25),
            "Heading 3": (11.5, 8, 4, 1.25),
            "Heading 4": (10.5, 6, 3, 1.25),
        }
    elif page_size == "b5":
        body_size, body_spacing = 10.0, 1.4
        specs = {
            "Heading 1": (15, 8, 4, 1.18),
            "Heading 2": (12, 9, 5, 1.22),
            "Heading 3": (10.5, 8, 3, 1.22),
            "Heading 4": (10, 6, 3, 1.22),
        }
    else:
        body_size, body_spacing = 9.5, 1.3
        specs = {
            "Heading 1": (14, 6, 3, 1.15),
            "Heading 2": (11.5, 8, 5, 1.2),
            "Heading 3": (10.5, 8, 3, 1.2),
            "Heading 4": (9.5, 6, 3, 1.2),
        }

    normal = doc.styles["Normal"]
    style_fonts(normal, BODY_CN, BODY_EN, body_size, False, BLACK)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = body_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(0.67)
    pf.widow_control = True

    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    rpr.append(lang)

    for name, (size, before, after, spacing) in specs.items():
        st = doc.styles[name]
        style_fonts(st, HEAD_CN, HEAD_EN, size, False, ACCENT_RGB)
        pf = st.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True
        pf.line_spacing = spacing
        if name == "Heading 1":
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        style_fonts(st, BODY_CN, BODY_EN, body_size, False, BLACK)
        st.paragraph_format.line_spacing = body_spacing
        st.paragraph_format.space_after = Pt(0)

    formula = doc.styles.add_style("Book Formula", 1)
    formula.base_style = doc.styles["Normal"]
    style_fonts(formula, BODY_CN, BODY_EN, body_size + 0.5, False, ACCENT_RGB)
    fpf = formula.paragraph_format
    fpf.first_line_indent = Cm(0)
    fpf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fpf.left_indent = Cm(1.0)
    fpf.right_indent = Cm(1.0)
    fpf.space_before = Pt(4)
    fpf.space_after = Pt(4)
    fpf.line_spacing = 1.2
    fpf.keep_together = True

    callout = doc.styles.add_style("Book Callout", 1)
    callout.base_style = doc.styles["Normal"]
    style_fonts(callout, BODY_CN, BODY_EN, body_size, False, BLACK)
    cpf = callout.paragraph_format
    cpf.first_line_indent = Cm(0)
    cpf.line_spacing = body_spacing + 0.05
    cpf.space_before = Pt(0)
    cpf.space_after = Pt(0)
    cpf.keep_together = True

    code = doc.styles.add_style("Book Code", 1)
    code.base_style = doc.styles["Normal"]
    style_fonts(code, BODY_CN, CODE_EN, 8.5, False, BLACK)
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.left_indent = Cm(0.3)
    code.paragraph_format.line_spacing = 1.15
    code.paragraph_format.space_before = Pt(6)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.keep_together = True


def setup_header(doc, header_title=None, header_author=None):
    header = doc.sections[0].header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_after = Pt(2)
    pf.tab_stops.clear_all()
    pf.tab_stops.add_tab_stop(Cm(text_width_cm(doc)), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(header_title or HEADER_TITLE)
    run_fonts(run, BODY_CN, BODY_EN, 8, color=MID_GRAY)
    p.add_run("\t")
    author = p.add_run(header_author or HEADER_AUTHOR)
    run_fonts(author, BODY_CN, BODY_EN, 7.5, color=RGBColor(0x9A, 0x9A, 0x9A))
    bottom_border(p, RULE_GRAY, 4, space=1)


def setup_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    prefix = p.add_run("— ")
    run_fonts(prefix, BODY_CN, BODY_EN, 9, color=MID_GRAY)
    field = add_page_field(p)
    run_fonts(field, BODY_CN, BODY_EN, 9, color=MID_GRAY)
    suffix = p.add_run(" —")
    run_fonts(suffix, BODY_CN, BODY_EN, 9, color=MID_GRAY)


def add_cover(doc, title, subtitle, description, edition):
    """Typographic A5 cover page with accent rules."""
    doc.sections[0].different_first_page_header_footer = True

    def centered(text, size, cn, en, color, before=0, after=0, spacing=1.2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = spacing
        run = p.add_run(text)
        run_fonts(run, cn, en, size, False, color)
        return p

    spacer = doc.add_paragraph()
    spacer.paragraph_format.first_line_indent = Cm(0)
    spacer.paragraph_format.space_before = Pt(112)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0

    centered(title, 27, HEAD_CN, HEAD_EN, ACCENT_RGB, after=8, spacing=1.1)
    centered(subtitle, 12.5, BODY_CN, BODY_EN, RGBColor(0x44, 0x44, 0x44),
             after=16, spacing=1.3)

    side = max((text_width_cm(doc) - 2.2) / 2, 0)
    rule = doc.add_paragraph()
    rule.paragraph_format.left_indent = Cm(side)
    rule.paragraph_format.right_indent = Cm(side)
    rule.paragraph_format.first_line_indent = Cm(0)
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(0)
    paragraph_borders(rule, {"bottom": (ACCENT, 12, 4)})

    centered(description, 9.5, BODY_CN, BODY_EN, GRAY, before=20, spacing=1.6)

    spacer2 = doc.add_paragraph()
    spacer2.paragraph_format.first_line_indent = Cm(0)
    spacer2.paragraph_format.space_before = Pt(180)
    spacer2.paragraph_format.space_after = Pt(0)
    spacer2.paragraph_format.line_spacing = 1.0

    centered(edition, 9, BODY_CN, BODY_EN, MID_GRAY, spacing=1.2)
    # Page break before the first chapter is set on that heading itself,
    # so no extra break paragraph can overflow onto a blank page.


# ---------------------------------------------------------------------------
# content helpers
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# inline sub/superscript
# ---------------------------------------------------------------------------
_SCRIPT_MARKS = ("_", "^")
_TOKEN_RE = re.compile(r"[A-Za-z0-9\u0391-\u03c9\u0370-\u03ff]+")


def _read_script(text, i, mark):
    """从 text[i]（'_' 或 '^'）读取一个上下标内容，返回 (content, end)。"""
    j = i + 1
    if j >= len(text):
        return None
    if text[j] == "{":
        k = text.find("}", j + 1)
        if k != -1 and "\n" not in text[j:k]:
            return text[j + 1:k], k + 1
        return None
    if text[j] == "(":
        depth = 0
        k = j
        while k < len(text):
            if text[k] == "(":
                depth += 1
            elif text[k] == ")":
                depth -= 1
                if depth == 0:
                    return text[j + 1:k], k + 1
            k += 1
        return None
    if text[j] == "\\":
        return None
    if mark == "_" and text[j] == "|":
        k = text.find("|", j + 1)
        if k > j + 1 and "\n" not in text[j:k]:
            return text[j:k + 1], k + 1
        return None
    if mark == "_" and re.match(r"[\u4e00-\u9fff]", text[j]):
        return text[j], j + 1
    if mark == "^":
        m = re.match(r"[+\-\u2212]?[A-Za-z0-9\u0391-\u03c9\u0370-\u03ff]+", text[j:])
    else:
        m = _TOKEN_RE.match(text[j:])
    if m:
        return m.group(0), j + m.end()
    return None



def _flatten_script(s):
    """上下标内容里若再嵌套标记（极少见），去掉嵌套的标记符，按纯文本排。"""
    out = []
    i = 0
    n = len(s)
    while i < n:
        if s[i] in _SCRIPT_MARKS:
            r = _read_script(s, i, s[i])
            if r:
                out.append(r[0])
                i = r[1]
                continue
        out.append(s[i])
        i += 1
    return "".join(out)


def add_rich_text(paragraph, text, size=None, color=None, bold=False,
                  cn=BODY_CN, en=BODY_EN):
    """写入一段文本，把 X_abc / X^abc / X_{..} / X^{..} / e^(..) 渲染成上下标。"""
    buf = []

    def flush():
        if buf:
            run = paragraph.add_run("".join(buf))
            run_fonts(run, cn, en, size, bold or None, color)
            buf.clear()

    i = 0
    n = len(text)
    while i < n:
        ch = text[i]
        if ch in _SCRIPT_MARKS:
            parsed = _read_script(text, i, ch)
            if parsed:
                content, end = parsed
                content = _flatten_script(content)
                flush()
                run = paragraph.add_run(content)
                run_fonts(run, cn, en, size, bold or None, color)
                if ch == "_":
                    run.font.subscript = True
                else:
                    run.font.superscript = True
                i = end
                continue
        buf.append(ch)
        i += 1
    flush()


def add_inline(paragraph, text, size=None, color=None, bold=False):
    text = smart_quotes(text)
    # Inline code backticks are typographic noise in print: keep the words.
    text = text.replace("`", "")
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            add_rich_text(paragraph, text[pos:m.start()], size, color, bold)
        add_rich_text(paragraph, m.group(1), size, color, True)
        pos = m.end()
    if pos < len(text):
        add_rich_text(paragraph, text[pos:], size, color, bold)


def smart_quotes(text):
    out = []
    opening = True
    for ch in text:
        if ch == '"':
            out.append("\u201c" if opening else "\u201d")
            opening = not opening
        else:
            out.append(ch)
    return "".join(out)


def add_title_rule(doc):
    width = text_width_cm(doc)
    side = max((width - 2.2) / 2, 0)
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(side)
    pf.right_indent = Cm(side)
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(2)
    pf.space_after = Pt(12)
    paragraph_borders(p, {"bottom": (ACCENT, 12, 6)})


def add_heading(doc, text, level):
    text = smart_quotes(text)
    level = min(level, 4)
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    size = {1: 14, 2: 11.5, 3: 10.5, 4: 9.5}[level]
    run_fonts(run, HEAD_CN, HEAD_EN, size, level == 2, ACCENT_RGB)
    if level == 1:
        add_title_rule(doc)
    elif level == 2:
        paragraph_borders(p, {"left": (ACCENT, 18, 6)})
    return p


def chapter_kicker(text):
    m = re.match(r"^★?\s*第\s*(\d+)\s*章", text)
    if m:
        return f"CHAPTER {int(m.group(1)):02d}"
    if text.startswith("序章"):
        return "PROLOGUE"
    if text.startswith("结语"):
        return "EPILOGUE"
    if text.startswith("延伸阅读"):
        return "FURTHER READING"
    if text.startswith("附录"):
        return "APPENDIX"
    return None


def add_chapter_kicker(doc, text):
    label = chapter_kicker(text)
    if not label:
        return None
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(8)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    run = p.add_run(label)
    run_fonts(run, HEAD_CN, HEAD_EN, 8, True, MID_GRAY)
    rpr = run._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "40")
    rpr.append(spacing)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    add_inline(p, text)
    return p


def add_formula(doc, text):
    p = doc.add_paragraph(style="Book Formula")
    add_rich_text(p, text.strip(), size=10, color=ACCENT_RGB)
    shade_paragraph(p, FORMULA_FILL)
    return p


def add_callout(doc, groups):
    total = len(groups)
    m = TITLE_RE.fullmatch(groups[0].strip()) if groups else None
    adv = bool(m and m.group(1).strip() == "进阶笔记")
    accent = ADV_ACCENT if adv else ACCENT
    fill = ADV_TINT if adv else CALLOUT_FILL
    for idx, text in enumerate(groups):
        p = doc.add_paragraph(style="Book Callout")
        pf = p.paragraph_format
        pf.left_indent = Cm(0.35)
        pf.right_indent = Cm(0.25)
        pf.space_before = Pt(2 if idx == 0 else 1)
        pf.space_after = Pt(4 if idx == total - 1 else 0)
        edges = {"left": (accent, 18, 6)}
        if idx == 0:
            edges["top"] = (accent, 6, 4)
        if idx == total - 1:
            edges["bottom"] = (accent, 6, 4)
        paragraph_borders(p, edges)
        shade_paragraph(p, fill)
        pf.keep_with_next = True
        m = TITLE_RE.fullmatch(text)
        if m:
            add_rich_text(p, m.group(1), size=10 if not adv else 9.5,
                          color=ADV_RGB if adv else ACCENT_RGB, bold=True,
                          cn=HEAD_CN, en=HEAD_EN)
        else:
            add_inline(p, text, size=9.5)
    return None


def add_list(doc, items, ordered):
    if ordered:
        # Manual numbering keeps each list starting at 1 and avoids Word's
        # document-wide list numbering continuing across chapters.
        chain = len(items) <= 3
        # 章末习题（≥4 题）用紧凑版式，减少最后几题被挤到单独一页
        compact = len(items) >= 4
        for idx, item in enumerate(items, 1):
            p = doc.add_paragraph(style="Normal")
            pf = p.paragraph_format
            pf.left_indent = Cm(0.74)
            pf.first_line_indent = Cm(-0.74)
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            if compact:
                pf.line_spacing = 1.15
            if chain and idx < len(items):
                pf.keep_with_next = True
            run = p.add_run(f"{idx}. ")
            run_fonts(run, BODY_CN, BODY_EN, 9 if compact else 9.5)
            add_inline(p, item, size=9 if compact else 9.5)
        return
    keep = len(items) <= 5
    last = len(items) - 1
    for idx, item in enumerate(items):
        p = doc.add_paragraph(style="List Bullet")
        pf = p.paragraph_format
        pf.left_indent = Cm(0.74)
        pf.first_line_indent = Cm(-0.37)
        if keep and idx < last:
            p.paragraph_format.keep_with_next = True
        add_inline(p, item)


def add_code(doc, code_lines):
    p = doc.add_paragraph(style="Book Code")
    for idx, line in enumerate(code_lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        run_fonts(run, BODY_CN, CODE_EN, 9)
    shade_paragraph(p, "F5F5F5")
    bottom_border(p, RULE_GRAY, 4)


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    bottom_border(p, RULE_GRAY, 4)


def split_table_row(raw: str) -> list[str]:
    """按未转义的竖线切分表格行，并把 \\| 还原成 |。"""
    s = raw.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    s = s.replace("\\|", "\x00")
    return [c.strip().replace("\x00", "|") for c in s.split("|")]


def add_table(doc, raw_rows, width):
    rows = []
    for raw in raw_rows:
        cells = split_table_row(raw)
        if cells and all(TABLE_SEP_RE.match(c) for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return 0
    ncols = len(rows[0])
    keep_table = len(rows) <= 6
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    table_borders(table)

    def text_units(text):
        return sum(2 if ord(ch) > 127 else 1 for ch in text)

    weights = []
    for j in range(ncols):
        longest = 0
        for row in rows:
            cell = row[j] if j < len(row) else ""
            longest = max(longest, text_units(cell))
        weights.append(max(4, longest))

    if ncols == 1:
        widths = [width]
    else:
        min_cm = 1.6 if ncols <= 5 else 1.4
        widths = [0.0] * ncols
        pending = list(range(ncols))
        remaining = width
        while pending:
            total = sum(weights[i] for i in pending)
            changed = False
            for i in list(pending):
                part = remaining * weights[i] / total
                if part < min_cm:
                    widths[i] = min_cm
                    remaining -= min_cm
                    pending.remove(i)
                    changed = True
            if not changed:
                for i in pending:
                    widths[i] = remaining * weights[i] / total
                break
    for j in range(ncols):
        table.columns[j].width = Cm(widths[j])

    # Alignment is a property of the column, not of a single cell: a text-heavy
    # column stays left-aligned even where one cell happens to be short.
    col_left = [
        j == 0 or max(
            text_units(row[j] if j < len(row) else "") for row in rows
        ) > 14
        for j in range(ncols)
    ]

    for r_idx, row in enumerate(rows):
        no_split_row(table.rows[r_idx])
        if r_idx == 0:
            mark_header_row(table.rows[r_idx])
        for c_idx in range(ncols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Cm(widths[c_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.space_before = Pt(1.0)
            pf.space_after = Pt(1.0)
            pf.line_spacing = 1.0
            p.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if col_left[c_idx]
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            text = row[c_idx] if c_idx < len(row) else ""
            if keep_table and r_idx < len(rows) - 1:
                pf.keep_with_next = True
            if r_idx == 0:
                add_inline(p, text, size=9, color=ACCENT_RGB, bold=True)
                shade_cell(cell, ACCENT_TINT)
                cell_bottom_border(cell, ACCENT, 8)
            else:
                add_inline(p, text, size=9)
                if r_idx % 2 == 1:
                    shade_cell(cell, ZEBRA_FILL)
    return 1


# ---------------------------------------------------------------------------
# markdown parser
# ---------------------------------------------------------------------------

def starts_block(line):
    s = line.strip()
    return bool(
        HEADING_RE.match(s)
        or s.startswith(">")
        or s.startswith("$$")
        or s.startswith("|")
        or s.startswith("```")
        or BULLET_RE.match(s)
        or NUMBER_RE.match(s)
        or IMAGE_RE.match(s)
        or s == "---"
        or s == "<!-- pagebreak -->"
    )


def add_figure(doc, caption, image_path):
    """Centered line-art figure with a small caption below."""
    width = min(text_width_cm(doc), 10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(8)
    pf.space_after = Pt(2)
    pf.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Cm(width))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = cap.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(12)
    pf.line_spacing = 1.2
    m = re.match(r"^(图\s*\d+(?:[-–]\d+)?)\s*(.*)$", caption)
    if m:
        lead = cap.add_run(m.group(1))
        run_fonts(lead, BODY_CN, BODY_EN, 9, True, ACCENT_RGB)
        add_rich_text(cap, (" " + m.group(2)) if m.group(2) else "", size=9, color=MID_GRAY)
    else:
        run = cap.add_run(caption)
        run_fonts(run, BODY_CN, BODY_EN, 9, color=MID_GRAY)


def add_toc(doc, entries, start_new_page=True):
    """Static table of contents with dot leaders and page numbers.

    The entries are wrapped in a real Word TOC field so that pressing F9 in
    Word can regenerate them; the static result keeps the PDF correct.
    """
    width = text_width_cm(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = title.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(18)
    pf.space_after = Pt(4)
    pf.page_break_before = start_new_page
    run = title.add_run("目录")
    run_fonts(run, HEAD_CN, HEAD_EN, 16, False, ACCENT_RGB)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.first_line_indent = Cm(0)
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(10)
    sub_run = sub.add_run("CONTENTS")
    run_fonts(sub_run, HEAD_CN, HEAD_EN, 7.5, True, MID_GRAY)
    rpr = sub_run._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "80")
    rpr.append(spacing)

    side = max((width - 2.2) / 2, 0)
    rule = doc.add_paragraph()
    rule.paragraph_format.left_indent = Cm(side)
    rule.paragraph_format.right_indent = Cm(side)
    rule.paragraph_format.first_line_indent = Cm(0)
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(14)
    paragraph_borders(rule, {"bottom": (ACCENT, 12, 6)})

    # TOC field: begin + instruction + separator; entries are cached result.
    start = doc.add_paragraph()
    spf = start.paragraph_format
    spf.first_line_indent = Cm(0)
    spf.space_before = Pt(0)
    spf.space_after = Pt(0)
    spf.line_spacing = Pt(1)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    start.add_run()._r.append(begin)
    instr_run = start.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    instr_run._r.append(instr)
    sep_run = start.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)

    for level, text, page in entries:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0) if level <= 1 else Cm(0.6)
        pf.space_before = Pt(2 if level <= 1 else 0)
        pf.space_after = Pt(0)
        pf.line_spacing = 0.98
        pf.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT,
                                  WD_TAB_LEADER.DOTS)
        if level <= 1:
            run = p.add_run(text)
            run_fonts(run, HEAD_CN, HEAD_EN, 10, True, ACCENT_RGB)
        else:
            run = p.add_run(text)
            run_fonts(run, BODY_CN, BODY_EN, 9, color=GRAY)
        p.add_run("\t")
        number = p.add_run(str(page))
        run_fonts(number, BODY_CN, BODY_EN, 9 if level > 1 else 10,
                  level <= 1, MID_GRAY if level > 1 else ACCENT_RGB)

    endp = doc.add_paragraph()
    epf = endp.paragraph_format
    epf.first_line_indent = Cm(0)
    epf.space_before = Pt(0)
    epf.space_after = Pt(0)
    epf.line_spacing = Pt(1)
    end_run = endp.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def build(md_path: Path, docx_path: Path, page_size: str = "a5",
          cover: bool = False, chapter_breaks: bool = False,
          toc_entries=None, asset_root: Path | None = None,
          title_text: str | None = None, subtitle_text: str | None = None,
          description_text: str | None = None, edition_text: str | None = None,
          header_title: str | None = None, header_author: str | None = None):
    doc = Document()
    setup_page(doc, page_size)
    setup_styles(doc, page_size)
    setup_header(doc, header_title, header_author)
    setup_footer(doc)
    if cover:
        add_cover(
            doc,
            title_text or BOOK_TITLE,
            subtitle_text or BOOK_SUBTITLE,
            description_text or BOOK_DESCRIPTION,
            edition_text or BOOK_EDITION,
        )
    if toc_entries:
        add_toc(doc, toc_entries, start_new_page=cover)

    width = text_width_cm(doc)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    title = None
    tables = 0
    last_block = None
    first_heading = True
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped == "<!-- pagebreak -->":
            doc.add_page_break()
            last_block = "pagebreak"
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            add_code(doc, code_lines)
            last_block = "code"
            continue

        m = HEADING_RE.match(stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and title is None:
                title = text
            kicker = add_chapter_kicker(doc, text) if level == 1 else None
            p = add_heading(doc, text, level)
            if level == 1:
                target = kicker if kicker is not None else p
                if first_heading:
                    if cover:
                        target.paragraph_format.page_break_before = True
                    first_heading = False
                elif chapter_breaks:
                    target.paragraph_format.page_break_before = True
            last_block = "heading"
            i += 1
            continue

        if stripped.startswith("$$"):
            m = FORMULA_RE.match(stripped)
            if m:
                add_formula(doc, m.group(1))
            else:
                add_formula(doc, stripped.strip("$").strip())
            last_block = "formula"
            i += 1
            continue

        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q = re.sub(r"^>\s?", "", lines[i].strip())
                quote.append(q)
                i += 1
            groups = []
            cur = []
            for q in quote:
                if q == "":
                    if cur:
                        groups.append("".join(cur))
                        cur = []
                else:
                    cur.append(q)
            if cur:
                groups.append("".join(cur))
            if groups:
                add_callout(doc, groups)
            last_block = "callout"
            continue

        if stripped.startswith("|"):
            nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
            is_table = bool(
                nxt.startswith("|")
                and all(TABLE_SEP_RE.match(c) for c in split_table_row(nxt))
            )
            if is_table:
                raw_rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    raw_rows.append(lines[i].strip())
                    i += 1
                tables += add_table(doc, raw_rows, width)
                last_block = "table"
                continue
            # 否则是正文（例如以 |V(z)|、|Γ| 开头的句子），按普通段落处理

        m = IMAGE_RE.match(stripped)
        if m:
            caption, rel = m.group(1).strip(), m.group(2).strip()
            candidates = []
            path = Path(rel)
            if path.is_absolute():
                candidates.append(path)
            else:
                candidates.append(md_path.parent / rel)
                if asset_root is not None:
                    candidates.append(asset_root / rel)
            image_path = next((c for c in candidates if c.exists()), None)
            if image_path is not None:
                add_figure(doc, caption, image_path)
            else:
                print(f"WARN missing image: {rel}")
            last_block = "figure"
            i += 1
            continue

        if BULLET_RE.match(stripped):
            items = []
            while i < len(lines) and BULLET_RE.match(lines[i].strip()):
                items.append(BULLET_RE.sub("", lines[i].strip(), count=1))
                i += 1
            add_list(doc, items, ordered=False)
            last_block = "list"
            continue

        if NUMBER_RE.match(stripped):
            items = []
            while i < len(lines) and NUMBER_RE.match(lines[i].strip()):
                items.append(NUMBER_RE.sub("", lines[i].strip(), count=1))
                i += 1
            add_list(doc, items, ordered=True)
            last_block = "list"
            continue

        if stripped == "---":
            add_rule(doc)
            last_block = "rule"
            i += 1
            continue

        para_lines = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not starts_block(lines[i]):
            para_lines.append(lines[i].strip())
            i += 1
        p = add_body(doc, "".join(para_lines))
        if last_block == "table":
            p.paragraph_format.space_before = Pt(8)
        if re.fullmatch(r"\*\*.+\*\*", stripped):
            p.paragraph_format.keep_with_next = True
        if re.fullmatch(r"自检.{0,4}：", p.text):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(8)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = ACCENT_RGB
        look = i
        while look < len(lines) and not lines[look].strip():
            look += 1
        if look < len(lines):
            nxt = lines[look].strip()
            if nxt.startswith("|") or BULLET_RE.match(nxt) or NUMBER_RE.match(nxt):
                p.paragraph_format.keep_with_next = True
        last_block = "body"

    doc.core_properties.title = title_text or title or md_path.stem
    doc.core_properties.subject = subtitle_text or BOOK_SUBTITLE
    doc.core_properties.author = header_author or HEADER_AUTHOR
    doc.core_properties.comments = f"Source: {md_path.name}"
    doc.save(docx_path)

    return {
        "title": title,
        "paragraphs": len(doc.paragraphs),
        "tables": tables,
        "page_size": page_size,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Markdown file")
    parser.add_argument("output", type=Path, help="DOCX output file")
    size = parser.add_mutually_exclusive_group()
    size.add_argument("--a4", action="store_true", help="use A4 instead of A5")
    size.add_argument("--b5", action="store_true", help="use B5 (176x250 mm)")
    parser.add_argument("--cover", action="store_true", help="add a cover page")
    parser.add_argument("--title", default=None, help="cover title")
    parser.add_argument("--subtitle", default=None, help="cover subtitle")
    parser.add_argument("--desc", default=None, help="cover description")
    parser.add_argument("--edition", default=None, help="cover edition line")
    parser.add_argument("--header", default=None, help="running header title")
    parser.add_argument("--author", default=None, help="running header author")
    args = parser.parse_args()
    page_size = "a4" if args.a4 else ("b5" if args.b5 else "a5")
    info = build(args.input, args.output, page_size,
                 cover=args.cover,
                 title_text=args.title, subtitle_text=args.subtitle,
                 description_text=args.desc, edition_text=args.edition,
                 header_title=args.header, header_author=args.author)
    print(f"OK title={info['title']!r} paragraphs={info['paragraphs']} "
          f"tables={info['tables']} page={info['page_size']} -> {args.output}")


if __name__ == "__main__":
    main()
