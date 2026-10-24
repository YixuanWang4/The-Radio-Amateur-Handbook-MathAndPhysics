#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Final integrity checks for the compiled book."""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document
from pdf_outline import extract

ROOT = Path(__file__).resolve().parent.parent
DOCX = ROOT / "build" / "业余无线电实用数学物理-书稿.docx"
PDF = ROOT / "build" / "业余无线电实用数学物理-书稿.pdf"


def main():
    lines = []
    d = Document(str(DOCX))
    residue = {"**": 0, "##": 0, "$$": 0, "`": 0}
    for p in d.paragraphs:
        for key in residue:
            if key in p.text:
                residue[key] += 1
    heads = {}
    for p in d.paragraphs:
        if p.style.name.startswith("Heading"):
            heads[p.style.name] = heads.get(p.style.name, 0) + 1
    images = len(d.inline_shapes)
    lines.append(f"paragraphs={len(d.paragraphs)} tables={len(d.tables)} images={images}")
    lines.append(f"headings={heads}")
    lines.append(f"markdown_residue={residue}")

    md_files = [p for p in sorted(ROOT.glob("*.md"))
                if re.match(r"^(0[1-9]|1[0-9]|2[0-5])-", p.name)]
    expected = 0
    for p in md_files:
        text = p.read_text(encoding="utf-8")
        for ln in text.splitlines():
            if ln.startswith("# ") or ln.startswith("## "):
                expected += 1
    entries = extract(PDF)
    lines.append(f"toc_entries={len(entries)} headings_in_sources={expected}")

    refs = set()
    for p in md_files:
        for m in re.finditer(r"!\[[^\]]*\]\((figures/fig-[0-9-]+\.png)\)",
                             p.read_text(encoding="utf-8")):
            refs.add(m.group(1))
    missing = [r for r in sorted(refs) if not (ROOT / r).exists()]
    lines.append(f"unique_figure_refs={len(refs)} missing={missing}")

    d_text = (ROOT / "23-附录D-习题答案.md").read_text(encoding="utf-8")
    lines.append(f"appendix_D_chapters={len(re.findall(r'^## ', d_text, re.M))}")

    out = "\n".join(lines)
    (ROOT / "build" / "final_check.txt").write_text(out, encoding="utf-8")
    print(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

def _pages():
    try:
        return len(PdfReader(str(PDF)).pages)
    except Exception:
        return None
